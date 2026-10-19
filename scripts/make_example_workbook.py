"""Generate examples/ProbabilityStepChart-Examples.xlsx (requires openpyxl).

Run with: python3 scripts/make_example_workbook.py

Sheets
  Start Here         - how to use the workbook with the add-in
  Make vs Buy        - the two example distributions from the specification
  Edge Cases         - inputs that exercise validation (sorting, duplicates, %, errors…)
  Dominance          - pairs with known FSD / SSD outcomes
  Reference Output   - the expected step coordinates and a native step chart for comparison
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "examples" / "ProbabilityStepChart-Examples.xlsx"
TITLE = Font(bold=True, size=12)
HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE = Font(italic=True, color="666666")


def block(ws, row, col, title, rows, note=None, prob_format=None, headers=("Outcome", "Probability")):
    """Write a titled two-column distribution; return the data range address (with header)."""
    ws.cell(row, col, title).font = TITLE
    if headers:
        for j, h in enumerate(headers):
            c = ws.cell(row + 1, col + j, h)
            c.font = HEAD
            c.fill = HEAD_FILL
        first = row + 2
    else:
        first = row + 1
    for i, (x, p) in enumerate(rows):
        ws.cell(first + i, col, x)
        cp = ws.cell(first + i, col + 1, p)
        if prob_format:
            cp.number_format = prob_format
    last = first + len(rows) - 1
    # Leave one blank row so a single-cell selection expands to just this block.
    if note:
        ws.cell(last + 2, col, note).font = NOTE
    start_row = row + 1 if headers else row + 1
    return f"{get_column_letter(col)}{start_row}:{get_column_letter(col + 1)}{last}"


def widths(ws, cols):
    for col, w in cols.items():
        ws.column_dimensions[col].width = w


def start_here(wb):
    ws = wb.active
    ws.title = "Start Here"
    lines = [
        ("Probability Step Chart – example workbook", Font(bold=True, size=14)),
        ("", None),
        ("1. Open Home > Step Chart (the add-in task pane).", None),
        ("2. On 'Make vs Buy', select A2:B5 (or just click any cell inside the table) and choose Use selection.", None),
        ("3. Check the preview, then choose Generate CDF coordinates. A table and an XY scatter step chart appear.", None),
        ("4. Select D2:E5 (Buy Component), choose Use selection, then Generate again: both CDFs share one chart.", None),
        ("5. Open the Analysis tab to compare the two distributions (expected value, variance, FSD, SSD).", None),
        ("   Tip: switch 'Outcomes are' to 'Lower is better (cost)' – for costs, Buy second-order dominates Make.", None),
        ("6. Change a probability in the source data, then Distributions > Regenerate all from source data.", None),
        ("", None),
        ("'Edge Cases' shows how validation handles unsorted, duplicate, zero-probability, negative and percentage inputs,", None),
        ("plus deliberately invalid inputs (sum ≠ 1, text in a number column).", None),
        ("'Dominance' contains pairs with known answers. 'Reference Output' shows the exact coordinates the add-in should produce.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(i, 1, text)
        if font:
            c.font = font
    widths(ws, {"A": 120})


def make_vs_buy(wb):
    ws = wb.create_sheet("Make vs Buy")
    block(ws, 1, 1, "Make Component", [(11, 0.4), (32, 0.3), (53, 0.3)])
    block(ws, 1, 4, "Buy Component", [(15, 0.4), (30, 0.3), (45, 0.3)])
    ws["A8"] = "Check: total probability"
    ws["A8"].font = NOTE
    ws["B8"] = "=SUM(B3:B5)"
    ws["D8"] = "Check: total probability"
    ws["D8"].font = NOTE
    ws["E8"] = "=SUM(E3:E5)"
    ws["A10"] = "Expected values: Make 29.9, Buy 28.5. Variances: Make 304.29, Buy 155.25."
    ws["A10"].font = NOTE
    widths(ws, {"A": 16, "B": 12, "C": 4, "D": 16, "E": 12})


def edge_cases(wb):
    ws = wb.create_sheet("Edge Cases")
    r = 1
    cases = [
        ("Unsorted outcomes", [(53, 0.3), (11, 0.4), (32, 0.3)], "Sorted automatically.", None),
        ("Duplicate outcomes", [(11, 0.2), (32, 0.3), (11, 0.2), (53, 0.3)], "11 appears twice; probabilities are combined (0.4).", None),
        ("Zero-probability outcome", [(11, 0.4), (20, 0), (32, 0.3), (53, 0.3)], "Kept for statistics; adds no step (warning).", None),
        ("Negative outcomes (profit/loss)", [(-20, 0.25), (0, 0.25), (50, 0.5)], "Automatic start X is a round number below -20.", None),
        ("Percent-formatted cells", [(11, 0.4), (32, 0.3), (53, 0.3)], "Cells show 40% but store 0.4.", "0%"),
        ("Percentage points (0-100)", [(11, 40), (32, 30), (53, 30)], "Auto-detected because they sum to 100.", None),
        ("Percent typed as text", [(11, "40%"), (32, "30%"), (53, "30%")], "Text such as '40%' is read as 0.4.", None),
        ("Single outcome", [(42, 1)], "One vertical jump from 0 to 100% at 42.", None),
        ("Ten equal outcomes (floating point)", [(i, 0.1) for i in range(1, 11)], "0.1 × 10 sums to 1 within tolerance; CDF ends at exactly 1.", None),
        ("INVALID: probabilities sum to 0.95", [(1, 0.5), (2, 0.45)], "Rejected with a clear message; nothing is written.", None),
        ("INVALID: text outcome", [("ten", 0.5), (20, 0.5)], "Rejected; the message names the cell.", None),
        ("INVALID: probability above 1", [(1, 1.2), (2, -0.2)], "Rejected: 1.2 > 1 and -0.2 < 0.", None),
    ]
    for title, rows, note, fmt in cases:
        block(ws, r, 1, title, rows, note=note, prob_format=fmt)
        r += len(rows) + 5
    widths(ws, {"A": 36, "B": 12})


def dominance(wb):
    ws = wb.create_sheet("Dominance")
    pairs = [
        (
            "First-order dominance: A FSD B",
            [(10, 0.5), (20, 0.5)],
            [(5, 0.5), (20, 0.5)],
            "A's CDF is never above B's (and is strictly below on [5, 10)). Every decision maker who prefers more chooses A.",
        ),
        (
            "Second-order only: sure 10 vs 50/50 of 0 or 20",
            [(10, 1)],
            [(0, 0.5), (20, 0.5)],
            "Same mean; the CDFs cross, so no FSD. A (the sure amount) SSD B: every risk-averse decision maker prefers it.",
        ),
        (
            "Higher mean and lower variance, but no SSD",
            [(-10, 0.01), (50, 0.99)],
            [(0, 0.5), (90, 0.5)],
            "A has EV 49.4 and variance 35.6; B has EV 45 and variance 2025, yet neither dominates (A's left tail at -10).",
        ),
        (
            "Identical distributions (different order)",
            [(11, 0.4), (32, 0.3), (53, 0.3)],
            [(53, 0.3), (11, 0.4), (32, 0.3)],
            "Reported as identical.",
        ),
    ]
    r = 1
    for title, a, b, note in pairs:
        ws.cell(r, 1, title).font = Font(bold=True, size=13)
        block(ws, r + 1, 1, "Distribution A", a)
        block(ws, r + 1, 4, "Distribution B", b)
        n = max(len(a), len(b))
        ws.cell(r + n + 4, 1, note).font = NOTE
        r += n + 7
    widths(ws, {"A": 16, "B": 12, "C": 4, "D": 16, "E": 12})


def reference_output(wb):
    ws = wb.create_sheet("Reference Output")
    make = [(0, 0), (11, 0), (11, 0.4), (32, 0.4), (32, 0.7), (53, 0.7), (53, 1)]
    buy = [(0, 0), (15, 0), (15, 0.4), (30, 0.4), (30, 0.7), (45, 0.7), (45, 1), (53, 1)]
    ws["A1"] = "Expected add-in output for Make vs Buy (start X = 0, final segment extended to the chart maximum 53)"
    ws["A1"].font = TITLE
    for col, title, pts in ((1, "Make Component", make), (4, "Buy Component", buy)):
        ws.cell(2, col, title).font = HEAD
        ws.cell(3, col, "X (Outcome)").font = HEAD
        ws.cell(3, col + 1, "Y (Cumulative Probability)").font = HEAD
        for i, (x, y) in enumerate(pts):
            ws.cell(4 + i, col, x)
            ws.cell(4 + i, col + 1, y)
    chart = ScatterChart()
    chart.title = "Cumulative Probability Distribution"
    chart.style = 2
    chart.scatterStyle = "line"
    chart.x_axis.title = "Outcome"
    chart.y_axis.title = "Cumulative Probability"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.y_axis.majorUnit = 0.1
    chart.y_axis.number_format = "0%"
    chart.x_axis.scaling.min = 0
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    for col, title, pts in ((1, "Make Component", make), (4, "Buy Component", buy)):
        xs = Reference(ws, min_col=col, min_row=4, max_row=3 + len(pts))
        ys = Reference(ws, min_col=col + 1, min_row=4, max_row=3 + len(pts))
        s = Series(ys, xs, title=title)
        s.marker.symbol = "none"
        s.smooth = False
        chart.series.append(s)
    chart.width = 18
    chart.height = 10
    ws.add_chart(chart, "H2")
    ws["A14"] = "This chart was drawn natively in the file to show what the add-in's chart should look like."
    ws["A14"].font = NOTE
    ws["A14"].alignment = Alignment(wrap_text=False)
    widths(ws, {"A": 14, "B": 26, "C": 4, "D": 14, "E": 26})


def main():
    wb = Workbook()
    start_here(wb)
    make_vs_buy(wb)
    edge_cases(wb)
    dominance(wb)
    reference_output(wb)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
